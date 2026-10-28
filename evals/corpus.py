"""Render the Markdown corpus to the PDF/DOCX files that users would upload.

The output is byte-for-byte deterministic, so re-running the eval finds the documents
already ingested (deduplication is by SHA-256) and does not pay for embeddings again.

Only a small Markdown subset is supported: "# " / "## " headings, paragraphs separated
by blank lines, and "- " list items.
"""

import io
import zipfile
from dataclasses import dataclass

FIXED_ZIP_TIME = (2025, 1, 1, 0, 0, 0)


@dataclass(frozen=True)
class Block:
    kind: str  # "h1", "h2" or "p"
    text: str


def parse_markdown(markdown: str) -> list[Block]:
    blocks = []
    for raw in markdown.split("\n\n"):
        lines = [line.strip() for line in raw.strip().splitlines()]
        if not lines:
            continue
        if lines[0].startswith("## "):
            blocks.append(Block("h2", lines[0][3:]))
        elif lines[0].startswith("# "):
            blocks.append(Block("h1", lines[0][2:]))
        elif all(line.startswith("- ") for line in lines):
            blocks.extend(Block("p", line) for line in lines)
        else:
            blocks.append(Block("p", " ".join(lines)))
    return blocks


def markdown_text(markdown: str) -> str:
    """The plain text a parser should get back from the rendered file."""
    return "\n\n".join(block.text for block in parse_markdown(markdown))


def render(markdown: str, fmt: str) -> bytes:
    if fmt == "pdf":
        return render_pdf(markdown)
    if fmt == "docx":
        return render_docx(markdown)
    raise ValueError(f"Unknown format: {fmt}")


def render_pdf(markdown: str) -> bytes:
    """A4 pages with wrapped text. Standard PDF fonts cover Latin text only."""
    # reportlab is a dev dependency: the eval is a development tool.
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import simpleSplit
    from reportlab.pdfgen import canvas

    styles = {"h1": ("Helvetica-Bold", 15), "h2": ("Helvetica-Bold", 12), "p": ("Helvetica", 10)}
    width, height = A4
    margin = 56
    buffer = io.BytesIO()
    # invariant=1: no creation date or random document id in the output.
    pdf = canvas.Canvas(buffer, pagesize=A4, invariant=1)
    y = height - margin
    for block in parse_markdown(markdown):
        font, size = styles[block.kind]
        leading = size * 1.4
        lines = simpleSplit(block.text, font, size, width - 2 * margin)
        if block.kind != "p" and y - leading * 3 < margin:
            pdf.showPage()  # keep a heading together with the text that follows
            y = height - margin
        for line in lines:
            if y < margin:
                pdf.showPage()
                y = height - margin
            pdf.setFont(font, size)
            pdf.drawString(margin, y, line)
            y -= leading
        y -= leading * 0.6
    pdf.save()
    return buffer.getvalue()


def render_docx(markdown: str) -> bytes:
    import docx

    document = docx.Document()
    for block in parse_markdown(markdown):
        if block.kind == "p":
            document.add_paragraph(block.text)
        else:
            document.add_heading(block.text, level=int(block.kind[1]))
    buffer = io.BytesIO()
    document.save(buffer)
    return _fix_zip_timestamps(buffer.getvalue())


def _fix_zip_timestamps(data: bytes) -> bytes:
    """ZIP entries carry the time they were written; pin it so the bytes are stable."""
    out = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(data)) as source,
        zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for item in source.infolist():
            info = zipfile.ZipInfo(item.filename, date_time=FIXED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(info, source.read(item.filename))
    return out.getvalue()
