"""Text extraction from uploaded files.

Parsers return a list of `Section`s that keep track of where the text came from
(PDF page number or DOCX heading), so chunks — and later, citations — can point back
to the source.
"""

import logging
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Literal

import docx
import pypdf
from docx.table import Table
from docx.text.paragraph import Paragraph

logger = logging.getLogger(__name__)

Kind = Literal["pdf", "docx"]

EXTENSIONS: dict[str, Kind] = {".pdf": "pdf", ".docx": "docx"}


class ParsingError(Exception):
    """The file cannot be turned into text; the message is shown to the user."""


@dataclass(frozen=True)
class Section:
    text: str
    page: int | None = None  # 1-based PDF page
    heading: str = ""  # nearest DOCX heading


@dataclass(frozen=True)
class ParsedDocument:
    sections: list[Section]
    page_count: int | None = None

    @property
    def text_length(self) -> int:
        return sum(len(s.text) for s in self.sections)


def detect_kind(head: bytes) -> Kind | None:
    """Identify the file type by its first bytes, not by the client-supplied name."""
    if head.startswith(b"%PDF-"):
        return "pdf"
    # DOCX is a ZIP archive; a cheap check for the ZIP signature, the parser validates
    # the rest (python-docx fails on ZIPs that are not Word documents).
    if head.startswith(b"PK\x03\x04"):
        return "docx"
    return None


def parse_document(path: Path, kind: Kind) -> ParsedDocument:
    parsed = _parse_pdf(path) if kind == "pdf" else _parse_docx(path)
    if parsed.text_length == 0:
        hint = " It may be a scanned document (OCR is not supported)." if kind == "pdf" else ""
        raise ParsingError(f"No extractable text found.{hint}")
    return parsed


# --- Normalisation -----------------------------------------------------------------

_HYPHEN_BREAK = re.compile(r"(\w)-\n(\w)")
_SPACES = re.compile(r"[ \t ]+")
_MANY_NEWLINES = re.compile(r"\n{3,}")


def normalize(text: str) -> str:
    # PostgreSQL text columns cannot store NUL bytes; some PDFs contain them.
    text = text.replace("\x00", "")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    # Words split by a hyphen at a line end ("infor-\nmation") are joined back.
    text = _HYPHEN_BREAK.sub(r"\1\2", text)
    text = _SPACES.sub(" ", text)
    text = "\n".join(line.strip() for line in text.split("\n"))
    return _MANY_NEWLINES.sub("\n\n", text).strip()


# --- PDF ---------------------------------------------------------------------------


def _parse_pdf(path: Path) -> ParsedDocument:
    try:
        reader = pypdf.PdfReader(path)
        if reader.is_encrypted:
            # Many PDFs are "encrypted" with an empty user password (copy protection).
            # decrypt() signals a wrong password by its return value, not an exception.
            if reader.decrypt("") == pypdf.PasswordType.NOT_DECRYPTED:
                raise ParsingError("The PDF is password-protected.")
        sections = []
        for number, page in enumerate(reader.pages, start=1):
            text = normalize(page.extract_text() or "")
            if text:
                sections.append(Section(text=text, page=number))
        return ParsedDocument(sections=sections, page_count=len(reader.pages))
    except ParsingError:
        raise
    except Exception as exc:  # pypdf raises many different errors on broken files
        logger.warning("PDF parsing failed for %s: %s", path, exc)
        raise ParsingError("The PDF file is damaged or not supported.") from exc


# --- DOCX --------------------------------------------------------------------------


def _parse_docx(path: Path) -> ParsedDocument:
    try:
        document = docx.Document(str(path))
    except Exception as exc:
        logger.warning("DOCX parsing failed for %s: %s", path, exc)
        raise ParsingError("The file is not a valid DOCX document.") from exc

    sections: list[Section] = []
    heading = ""
    buffer: list[str] = []

    def flush() -> None:
        text = normalize("\n\n".join(buffer))
        if text:
            sections.append(Section(text=text, heading=heading))
        buffer.clear()

    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if not text:
                continue
            if _is_heading(block):
                # Text before the next heading belongs to the previous one.
                flush()
                heading = text
                buffer.append(text)  # keep the heading in the text: it carries meaning
            else:
                buffer.append(text)
        elif isinstance(block, Table):
            buffer.append(_table_text(block))
    flush()
    # DOCX has no fixed pages: page numbers depend on rendering.
    return ParsedDocument(sections=sections, page_count=None)


def _is_heading(paragraph: Paragraph) -> bool:
    style = (paragraph.style.name if paragraph.style is not None else "") or ""
    return style.startswith("Heading") or style == "Title"


def _table_text(table: Table) -> str:
    rows = []
    for row in table.rows:
        cells = [cell.text.strip() for cell in row.cells]
        # Merged cells are repeated by python-docx; drop consecutive duplicates.
        deduped = [c for i, c in enumerate(cells) if i == 0 or c != cells[i - 1]]
        rows.append(" | ".join(deduped))
    return "\n".join(rows)
