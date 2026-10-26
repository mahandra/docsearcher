"""Build small PDF/DOCX files on the fly, so no binary fixtures live in the repo."""

import io

import docx
import pypdf
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas


def make_pdf(pages: list[list[str]]) -> bytes:
    """One inner list per page, one string per line. An empty list gives a blank page."""
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    for lines in pages:
        y = 800
        for line in lines:
            pdf.drawString(50, y, line)
            y -= 16
        pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def encrypt_pdf(data: bytes, password: str) -> bytes:
    writer = pypdf.PdfWriter(clone_from=pypdf.PdfReader(io.BytesIO(data)))
    writer.encrypt(user_password=password, owner_password=password + "-owner")
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def make_docx(blocks: list[tuple[str, object]]) -> bytes:
    """Blocks: ("h1"|"h2", text), ("p", text) or ("table", [[cell, ...], ...])."""
    document = docx.Document()
    for kind, value in blocks:
        if kind in ("h1", "h2"):
            document.add_heading(value, level=int(kind[1]))
        elif kind == "p":
            document.add_paragraph(value)
        elif kind == "table":
            table = document.add_table(rows=len(value), cols=len(value[0]))
            for r, row in enumerate(value):
                for c, cell in enumerate(row):
                    table.cell(r, c).text = cell
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
